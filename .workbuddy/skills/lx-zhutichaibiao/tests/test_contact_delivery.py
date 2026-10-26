"""Contact delivery regressions use real workbooks and ZIPs, never live services."""
import io
import json
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import split_by_zhuti as split
from lx_shujuku import build_mabiao_mapping


ROWS = [
    {'operator': '共同主体', 'brand': '品牌甲', 'city': '甲市', 'contact_person': '甲'},
    {'operator': '共同主体', 'brand': '品牌乙', 'city': '乙市', 'contact_person': '乙'},
]


def mapping(rows=None):
    result = build_mabiao_mapping(ROWS if rows is None else rows)
    result['_query_audit'] = {'is_complete': True, 'possible_truncation': False,
                              'row_count': len(ROWS if rows is None else rows)}
    return result


class ContactDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / '输入' / 'source.xlsx'
        self.source.parent.mkdir()
        self.config = {'项目根目录': str(self.root), '工作目录': str(self.root),
                       '默认对接人': ['甲'], '默认': {}, '特定配置': []}

    def workbook(self, rows=None, headers=None, keep=False):
        wb = Workbook()
        ws = wb.active
        ws.title = '数据'
        ws.append(headers or ['品牌名称', '城市', '编号'])
        for row in rows or [('品牌甲', '甲市', 1), ('品牌乙', '乙市', 2)]:
            ws.append(row)
        ws['A1'].font = Font(bold=True)
        ws.column_dimensions['A'].width = 24
        if keep:
            wb.create_sheet('保留信息').append(['其他人的完整明细'])
        wb.save(self.source)
        wb.close()

    def run_split(self, person='all', mode=1, keep=False, data=None):
        argv = ['-m', str(mode), '-p', person, '--confirmed']
        if keep:
            argv += ['-k', '保留信息']
        with patch.object(split, 'load_config', return_value=self.config), patch.object(
            split, 'load_company_mabiao', return_value=mapping() if data is None else data
        ), redirect_stdout(io.StringIO()):
            try:
                return split.main(argv)
            except Exception as exc:
                self.fail(f'CLI must report a failed task without losing input: {exc}')

    def packages(self):
        return sorted((self.root / '输出').rglob('*.zip'))

    def contents(self, package):
        records = []
        with zipfile.ZipFile(package) as zf:
            self.assertIsNone(zf.testzip())
            for name in zf.namelist():
                self.assertNotEqual(name, 'source.xlsx')
                wb = load_workbook(io.BytesIO(zf.read(name)), data_only=False)
                self.assertEqual(wb.sheetnames, ['数据'])
                records.extend(wb['数据'].iter_rows(min_row=2, values_only=True))
                self.assertTrue(wb['数据']['A1'].font.bold)
                self.assertEqual(wb['数据'].column_dimensions['A'].width, 24)
                wb.close()
        return records

    def test_shared_operator_is_split_before_each_contact_is_packaged(self):
        self.workbook(keep=True)
        self.assertEqual(self.run_split(keep=True), 0)
        packages = self.packages()
        self.assertEqual({p.name for p in packages}, {'甲.zip', '乙.zip'})
        self.assertEqual(self.contents(next(p for p in packages if p.name == '甲.zip')),
                         [('品牌甲', '甲市', 1)])
        self.assertEqual(self.contents(next(p for p in packages if p.name == '乙.zip')),
                         [('品牌乙', '乙市', 2)])
        self.assertTrue(list(self.root.rglob('*.xlsx')), 'local source remains recoverable')

    def test_duplicate_source_rows_survive_but_duplicate_mapping_does_not_expand_twice(self):
        self.workbook([('品牌甲', '甲市', 1), ('品牌甲', '甲市', 1)])
        self.assertEqual(self.run_split('甲', data=mapping([ROWS[0], ROWS[0]])), 0)
        self.assertEqual(self.contents(self.packages()[0]),
                         [('品牌甲', '甲市', 1), ('品牌甲', '甲市', 1)])

    def test_filtered_and_unmatched_rows_are_not_conflated_or_archived(self):
        self.workbook([('品牌甲', '甲市', 1), ('品牌乙', '乙市', 2), ('未知', '未知市', 3)])
        original = self.source.read_bytes()
        self.assertEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists(), 'unmatched input must remain available')
        self.assertEqual(self.source.read_bytes(), original)
        audits = [json.loads(p.read_text()) for p in (self.root / '输出').rglob('audit.json')]
        self.assertEqual(len(audits), 1)
        self.assertEqual(audits[0]['counts'],
                         {'source': 3, 'target': 1, 'filtered': 1, 'unmatched': 1,
                          'expected_relations': 1, 'actual_relations': 1})
        self.assertFalse(audits[0]['complete'])

    def test_missing_exact_mapping_and_unknown_contact_fail_without_moving_source(self):
        self.workbook()
        for data, person in [(dict(mapping(), operator_brand_rows=[]), '甲'), (mapping(), '不存在')]:
            with self.subTest(person=person):
                self.assertNotEqual(self.run_split(person, data=data), 0)
                self.assertTrue(self.source.exists())
                self.assertEqual(self.packages(), [])

    def test_source_only_mode_cannot_silently_fallback_to_city(self):
        self.workbook([('甲市', 1)], headers=['城市', '编号'])
        self.assertNotEqual(self.run_split('甲', mode=1), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(self.packages(), [])

    def test_city_mode_can_legitimately_expand_to_two_contacts(self):
        self.workbook([('共同市', 1)], headers=['城市', '编号'])
        data = mapping([dict(ROWS[0], city='共同市'), dict(ROWS[1], city='共同市')])
        self.assertEqual(self.run_split(mode=2, data=data), 0)
        self.assertEqual({p.name for p in self.packages()}, {'甲.zip', '乙.zip'})
        for package in self.packages():
            self.assertEqual(self.contents(package), [('共同市', 1)])

    def test_formula_dependency_fails_without_moving_input(self):
        self.workbook([('品牌甲', '甲市', "='保留信息'!A1")], keep=True)
        self.assertNotEqual(self.run_split('甲', keep=True), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(self.packages(), [])

    def test_local_report_failure_prevents_source_archive(self):
        self.workbook([('品牌甲', '甲市', 1)])
        original = self.source.read_bytes()
        with patch.object(split, 'write_processing_log', side_effect=OSError('log failed')):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertEqual(self.source.read_bytes(), original)

    def test_readback_rejects_style_loss_before_archive(self):
        self.workbook([('品牌甲', '甲市', 1)])
        write = split.write_delivery_files

        def corrupt_style(plan, task_dir):
            files = write(plan, task_dir)
            path = next(iter(files.values()))
            wb = load_workbook(path)
            wb['数据']['A1'].font = Font(bold=False)
            wb.save(path)
            wb.close()
            return files

        with patch.object(split, 'write_delivery_files', side_effect=corrupt_style):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])

    def test_brand_mode_is_scoped_by_available_city_and_pure_brand_is_local_only(self):
        self.workbook([('同品牌', '甲市', 1), ('同品牌', '乙市', 2)])
        data = mapping([dict(ROWS[0], brand='同品牌'), dict(ROWS[1], brand='同品牌')])
        self.assertEqual(self.run_split('甲', mode=3, data=data), 0)
        self.assertEqual(self.contents(self.packages()[0]), [('同品牌', '甲市', 1)])
        self.workbook()
        self.assertEqual(self.run_split(mode=4), 0)
        receipts = [json.loads(p.read_text()) for p in (self.root / '输出').rglob('ready.json')]
        local = [p for p in receipts if p['delivery_scope'] == 'local-only']
        self.assertEqual(len(local), 1)
        self.assertEqual(local[0]['packages'], ['本地品牌包/local-brand.zip'])

    def test_no_target_or_only_keep_sheet_never_moves_source(self):
        self.workbook([('品牌乙', '乙市', 2)])
        self.assertEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(self.packages(), [])
        self.workbook(keep=True)
        self.config['默认']['处理sheet'] = ['保留信息']
        self.assertNotEqual(self.run_split('甲', keep=True), 0)
        self.assertTrue(self.source.exists())

    def test_two_runs_do_not_overwrite_and_successful_archive_has_same_bytes(self):
        self.workbook([('品牌甲', '甲市', 1)])
        original = self.source.read_bytes()
        self.assertEqual(self.run_split('甲'), 0)
        first = {str(p): p.read_bytes() for p in self.packages()}
        self.assertFalse(self.source.exists())
        archived = list((self.root / '原表存档').rglob('source.xlsx'))
        self.assertEqual([p.read_bytes() for p in archived], [original])
        self.workbook([('品牌甲', '甲市', 9)])
        self.assertEqual(self.run_split('甲'), 0)
        self.assertEqual(len(self.packages()), 2)
        for path, data in first.items():
            self.assertEqual(Path(path).read_bytes(), data)

    def test_zip_extra_member_is_rejected_and_intermediates_retained(self):
        self.workbook([('品牌甲', '甲市', 1)])
        write = zipfile.ZipFile.write

        def extra(zf, *args, **kwargs):
            result = write(zf, *args, **kwargs)
            zf.writestr('unexpected.txt', 'must-not-deliver')
            return result

        with patch.object(zipfile.ZipFile, 'write', new=extra):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertTrue(list((self.root / '输出').rglob('共同主体.xlsx')))
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])

    def test_archive_failure_does_not_report_ready_or_delete_source(self):
        self.workbook([('品牌甲', '甲市', 1)])
        with patch.object(split.shutil, 'move', side_effect=OSError('archive failed')):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])
        failure = json.loads(next((self.root / '输出').rglob('failure.json')).read_text())
        self.assertEqual(failure['phase'], 'archive')

    def test_source_change_during_generation_is_not_archived_or_marked_ready(self):
        self.workbook([('品牌甲', '甲市', 1)])
        write = split.write_delivery_files

        def change_source(plan, task_dir):
            files = write(plan, task_dir)
            with self.source.open('ab') as out:
                out.write(b'changed')
            return files

        with patch.object(split, 'write_delivery_files', side_effect=change_source):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])

    def test_multiple_sheets_with_merged_headers_and_hidden_data(self):
        self.workbook([('品牌甲', '甲市', 1)])
        wb = load_workbook(self.source)
        ws = wb['数据']
        ws.insert_rows(1)
        ws['A1'] = '本次填写表'
        ws.merge_cells('A1:C1')
        wb.create_sheet('隐藏资料').sheet_state = 'hidden'
        wb.save(self.source)
        wb.close()
        self.config['默认']['表头行数'] = {'数据': 2}
        self.assertEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists(), 'unprocessed hidden sheet must retain source')
        with zipfile.ZipFile(self.packages()[0]) as zf:
            wb = load_workbook(io.BytesIO(zf.read(zf.namelist()[0])))
            self.assertEqual(wb.sheetnames, ['数据'])
            self.assertEqual(str(next(iter(wb['数据'].merged_cells.ranges))), 'A1:C1')
            self.assertEqual(list(wb['数据'].values)[2], ('品牌甲', '甲市', 1))
            wb.close()

    def test_sanitized_operator_name_collision_is_blocked(self):
        self.workbook([('品牌甲', '甲市', 1)])
        data = mapping([dict(ROWS[0], operator='主体/一'), dict(ROWS[0], operator='主体:一')])
        self.assertNotEqual(self.run_split('甲', data=data), 0)
        self.assertEqual(self.packages(), [])
        self.assertTrue(self.source.exists())

    def test_injected_extra_sheet_fails_independent_readback(self):
        self.workbook([('品牌甲', '甲市', 1)])
        write = split.write_delivery_files

        def add_sheet(plan, task_dir):
            files = write(plan, task_dir)
            path = next(iter(files.values()))
            wb = load_workbook(path)
            wb.create_sheet('其他人').append(['secret-in-fixture'])
            wb.save(path)
            wb.close()
            return files

        with patch.object(split, 'write_delivery_files', side_effect=add_sheet):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])

    def test_empty_input_directory_does_not_load_mapping(self):
        with patch.object(split, 'load_config', return_value=self.config), patch.object(
            split, 'load_company_mabiao', side_effect=AssertionError('no data to query')
        ), redirect_stdout(io.StringIO()):
            self.assertEqual(split.main(['-m', '1', '-p', '甲', '--confirmed']), 0)
        self.assertFalse((self.root / '输出').exists())

    def test_header_and_blank_row_are_accounted_for_outside_source_counts(self):
        self.workbook([('品牌甲', '甲市', 1), (None, None, None), ('品牌甲', '甲市', 2)])
        self.assertEqual(self.run_split('甲'), 0)
        audit = json.loads(next((self.root / '输出').rglob('audit.json')).read_text())
        self.assertEqual(audit['counts']['source'], 2)
        self.assertEqual(audit.get('excluded_rows'), [
            {'sheet': '数据', 'rows': [1, 1], 'reason': 'header'},
            {'sheet': '数据', 'rows': [3, 3], 'reason': 'blank'}])

    def test_data_area_merge_is_not_silently_lost(self):
        self.workbook([('品牌甲', '甲市', 1), ('品牌甲', '甲市', 2)])
        wb = load_workbook(self.source)
        wb['数据'].merge_cells('C2:C3')
        wb.save(self.source)
        wb.close()
        self.assertNotEqual(self.run_split('甲'), 0)
        self.assertTrue(self.source.exists())
        self.assertEqual(self.packages(), [])

    def test_receipt_failure_after_move_reports_actual_archive_without_rollback(self):
        self.workbook([('品牌甲', '甲市', 1)])
        original = self.source.read_bytes()
        open_path = Path.open

        def fail_receipt(path, *args, **kwargs):
            if path.name == 'ready.json':
                raise OSError('receipt failed after archive')
            return open_path(path, *args, **kwargs)

        with patch.object(Path, 'open', new=fail_receipt):
            self.assertNotEqual(self.run_split('甲'), 0)
        self.assertFalse(self.source.exists())
        failure = json.loads(next((self.root / '输出').rglob('failure.json')).read_text())
        self.assertEqual(failure['phase'], 'receipt')
        self.assertFalse(failure['source_present'])
        self.assertTrue(failure['archive_present'])
        self.assertEqual((self.root / failure['archive']).read_bytes(), original)
        self.assertEqual(list((self.root / '输出').rglob('ready.json')), [])


class CompleteMappingTests(unittest.TestCase):
    def test_complete_loader_uses_catalog_key_and_retains_evidence(self):
        from lx_shujuku.schema import SchemaCatalog

        row = {'id': 1, 'operator_entity': '共同主体', 'brand_name': '品牌甲',
               'city_name': '甲市', 'contact_person': '甲',
               'is_opening': '1', 'open_city_date': None}

        class Client:
            schema = SchemaCatalog.from_skill_root(Path(__file__).resolve().parents[2] / 'lx_shujuku')
            max_limit = 1000

            def execute_full_audited(client, sql, **kwargs):
                self.assertIn('ORDER BY id', sql)
                self.assertNotIn('LIMIT', sql)
                return {'rows': [row], 'is_complete': True, 'possible_truncation': False,
                        'row_count': 1, 'returned_rows': 1, 'total_rows': 1,
                        'result_sha256': 'test-hash',
                        'verification': {'passed': True, 'count_before': 1, 'count_after': 1}}

        with patch.object(split, 'create_client', return_value=Client()):
            result = split.load_company_mabiao()
        self.assertEqual(result['operator_brand_rows'][0]['contact_person'], '甲')
        self.assertTrue(result['_query_audit']['is_complete'])

    def test_loader_rejects_audited_incomplete_response(self):
        class Schema:
            def columns(self, name):
                return [{'field': 'id', 'key': 'PRI'}] + [
                    {'field': name, 'key': ''} for name in
                    ['operator_entity', 'brand_name', 'city_name', 'contact_person',
                     'is_opening', 'open_city_date']]

        class Client:
            schema = Schema()
            max_limit = 1000

            def execute_full_audited(self, sql, **kwargs):
                return {'rows': ROWS, 'is_complete': False, 'possible_truncation': True,
                        'row_count': 2, 'total_rows': 3, 'returned_rows': 2,
                        'verification': {'passed': False, 'count_before': 3, 'count_after': 3}}

            def load_mabiao_mapping(self):
                return mapping()

        with patch.object(split, 'create_client', return_value=Client()):
            with self.assertRaises(RuntimeError):
                split.load_company_mabiao()
